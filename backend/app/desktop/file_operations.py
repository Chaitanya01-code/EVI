from __future__ import annotations

import fnmatch
import hashlib
import logging
import os
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from app.desktop.path_guard import safe_path


logger = logging.getLogger(__name__)
_CODE_EXTENSIONS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".cs",
    ".go", ".rs", ".html", ".css", ".xml", ".yaml", ".yml", ".toml", ".ini",
    ".sql", ".sh", ".ps1", ".bat", ".md",
}
_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log"} | _CODE_EXTENSIONS
_MAX_BULK_ITEMS = 200


def _error(reason: str, **extra: Any) -> dict[str, Any]:
    return {"ok": False, "success": False, "verified": False, "reason": reason, **extra}


def _guard(path: str | os.PathLike[str]) -> tuple[Path | None, dict[str, Any] | None]:
    resolved, error = safe_path(path)
    return resolved, error


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _tree_hash(path: Path) -> str:
    digest = hashlib.sha256()
    for item in sorted(path.rglob("*"), key=lambda value: str(value).casefold()):
        guarded, error = _guard(item)
        if error or guarded is None:
            raise ValueError("denied_path")
        digest.update(item.relative_to(path).as_posix().encode("utf-8"))
        if item.is_file():
            digest.update(_hash(item).encode("ascii"))
    return digest.hexdigest()


def _validate_tree(path: Path) -> dict[str, Any] | None:
    if not path.is_dir():
        return None
    for index, item in enumerate(path.rglob("*"), 1):
        if index > _MAX_BULK_ITEMS:
            return _error("limit_exceeded", limit=_MAX_BULK_ITEMS)
        _resolved, error = _guard(item)
        if error:
            return error
    return None


def _walk_files(root: Path) -> Iterable[Path]:
    for current, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = [name for name in dirs if name.casefold() != ".git"]
        for name in files:
            candidate, error = _guard(Path(current) / name)
            if candidate is not None and error is None and candidate.is_file():
                yield candidate


def list_dir(path: str, depth: int = 1) -> dict[str, Any]:
    root, error = _guard(path)
    if error:
        return error
    if not root.exists() or not root.is_dir():
        return _error("not_a_directory")
    depth = max(1, min(int(depth), 10))
    rows = []
    for current, dirs, files in os.walk(root, followlinks=False):
        current_path = Path(current)
        level = len(current_path.relative_to(root).parts)
        dirs[:] = [
            name for name in dirs
            if name.casefold() != ".git" and _guard(current_path / name)[1] is None
        ]
        if level >= depth:
            dirs[:] = []
        for name in sorted(dirs + files, key=str.casefold):
            item_path, denied = _guard(current_path / name)
            if denied or item_path is None:
                continue
            try:
                stat = item_path.stat()
                rows.append({
                    "name": name,
                    "path": str(item_path),
                    "type": "directory" if item_path.is_dir() else "file",
                    "size": stat.st_size if item_path.is_file() else None,
                })
            except OSError:
                logger.warning("Unable to inspect directory entry %s", item_path, exc_info=True)
    return {"ok": True, "success": True, "verified": True, "path": str(root), "items": rows}


def search_files(
    root: str,
    name_glob: str = "*",
    extension: str = "",
    modified_after: str = "",
    limit: int = 200,
) -> dict[str, Any]:
    base, error = _guard(root)
    if error:
        return error
    if not base.exists() or not base.is_dir():
        return _error("not_a_directory")
    try:
        after = datetime.fromisoformat(modified_after).timestamp() if modified_after else None
    except ValueError:
        return _error("invalid_modified_after")
    extension = extension.casefold()
    if extension and not extension.startswith("."):
        extension = "." + extension
    limit = max(1, min(int(limit), _MAX_BULK_ITEMS))
    matches = []
    for path in _walk_files(base):
        try:
            if not fnmatch.fnmatch(path.name.casefold(), name_glob.casefold()):
                continue
            if extension and path.suffix.casefold() != extension:
                continue
            stat = path.stat()
            if after is not None and stat.st_mtime <= after:
                continue
            matches.append({"path": str(path), "name": path.name, "size": stat.st_size, "modified": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat()})
            if len(matches) >= limit:
                break
        except OSError:
            logger.warning("Unable to inspect search result %s", path, exc_info=True)
    return {"ok": True, "success": True, "verified": True, "matches": matches, "count": len(matches), "limit": limit}


def search_in_files(root: str, text: str, extensions: list[str] | None = None, limit: int = 200) -> dict[str, Any]:
    base, error = _guard(root)
    if error:
        return error
    if not base.exists() or not base.is_dir():
        return _error("not_a_directory")
    allowed = {ext.casefold() if ext.startswith(".") else "." + ext.casefold() for ext in extensions or []}
    limit = max(1, min(int(limit), _MAX_BULK_ITEMS))
    matches = []
    for path in _walk_files(base):
        if path.suffix.casefold() not in (_TEXT_EXTENSIONS if not allowed else allowed):
            continue
        try:
            with path.open("r", encoding="utf-8", errors="replace") as source:
                for line_number, line in enumerate(source, 1):
                    if text.casefold() in line.casefold():
                        matches.append({"path": str(path), "line": line_number, "text": line.strip()[:500]})
                        if len(matches) >= limit:
                            return {"ok": True, "success": True, "verified": True, "matches": matches, "count": len(matches), "limit": limit}
        except OSError:
            logger.warning("Unable to search file contents in %s", path, exc_info=True)
    return {"ok": True, "success": True, "verified": True, "matches": matches, "count": len(matches), "limit": limit}


def file_info(path: str) -> dict[str, Any]:
    resolved, error = _guard(path)
    if error:
        return error
    try:
        stat = resolved.stat()
        return {
            "ok": True, "success": True, "verified": True, "path": str(resolved),
            "size": stat.st_size, "created": datetime.fromtimestamp(stat.st_ctime).astimezone().isoformat(),
            "modified": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(),
            "type": "directory" if resolved.is_dir() else resolved.suffix.lstrip(".").lower() or "file",
        }
    except OSError as error:
        return _error("not_found", message=str(error))


def read_file(path: str, max_chars: int = 20000) -> dict[str, Any]:
    resolved, error = _guard(path)
    if error:
        return error
    if not resolved.is_file():
        return _error("not_a_file")
    extension = resolved.suffix.casefold()
    max_chars = max(1, min(int(max_chars), 20000))
    try:
        if extension in _TEXT_EXTENSIONS:
            text = resolved.read_text(encoding="utf-8", errors="replace")
        elif extension == ".docx":
            from docx import Document
            doc = Document(str(resolved))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        elif extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(resolved))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(str(resolved), read_only=True, data_only=True)
            chunks = []
            try:
                for sheet in workbook.worksheets[:5]:
                    chunks.append(f"[{sheet.title}]")
                    for row in sheet.iter_rows(max_row=1000, values_only=True):
                        chunks.append("\t".join("" if value is None else str(value) for value in row))
            finally:
                workbook.close()
            text = "\n".join(chunks)
        elif extension == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(resolved))
            text = "\n".join(
                shape.text for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
            )
        else:
            return {"ok": True, "success": True, "verified": True, "readable": False, "reason": "unsupported_type", "suggestion": "Use open_file to open this type."}
        return {"ok": True, "success": True, "verified": True, "readable": True, "text": text[:max_chars], "truncated": len(text) > max_chars}
    except Exception as error:
        logger.warning("Unable to read document %s", resolved, exc_info=True)
        return _error("read_failed", message=str(error))


def create_folder(path: str) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if target.exists():
        return _error("exists")
    try:
        target.mkdir(parents=True, exist_ok=False)
        verified = target.is_dir()
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"exists": verified}}
    except OSError as error:
        return _error("write_failed", message=str(error))


def create_file(path: str, content: str) -> dict[str, Any]:
    return write_file(path, content, mode="create")


def write_file(path: str, content: str, mode: str = "overwrite", overwrite: bool = False) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if mode not in {"overwrite", "append", "create"}:
        return _error("invalid_mode")
    if target.exists() and mode == "create":
        return _error("exists")
    if target.exists() and mode == "overwrite" and not overwrite:
        return _error("exists", message="Set overwrite=True to replace this file.")
    if target.exists() and mode == "append":
        original_size = target.stat().st_size
    else:
        original_size = 0
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("a" if mode == "append" else "x" if mode == "create" or (mode == "overwrite" and not target.exists()) else "w", encoding="utf-8", newline="") as output:
            output.write(content)
        actual = target.read_text(encoding="utf-8")
        expected = content if mode != "append" else None
        verified = target.is_file() and (actual == expected if expected is not None else actual.endswith(content) and target.stat().st_size >= original_size)
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"size": target.stat().st_size, "sha256": _hash(target)}}
    except FileExistsError:
        return _error("exists")
    except OSError as error:
        return _error("write_failed", message=str(error))


def create_docx(path: str, title: str, paragraphs: list[str]) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if target.exists():
        return _error("exists")
    try:
        from docx import Document
        target.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        if title:
            document.add_heading(title, 0)
        for paragraph in paragraphs:
            document.add_paragraph(str(paragraph))
        document.save(str(target))
        verified = target.is_file() and target.stat().st_size > 0
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"size": target.stat().st_size if verified else 0}}
    except (OSError, ImportError) as error:
        return _error("write_failed", message=str(error))


def create_xlsx(path: str, rows: list[list[Any]]) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if target.exists():
        return _error("exists")
    try:
        from openpyxl import Workbook, load_workbook
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        workbook.save(str(target))
        check = load_workbook(str(target), read_only=True, data_only=True)
        try:
            verified = list(check.active.iter_rows(values_only=True)) == [tuple(row) for row in rows]
        finally:
            check.close()
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"size": target.stat().st_size}}
    except (OSError, ImportError) as error:
        return _error("write_failed", message=str(error))


def _destination(src: Path, dst_value: str) -> tuple[Path | None, dict[str, Any] | None]:
    dst, error = _guard(dst_value)
    if error:
        return None, error
    if dst.exists() and dst.is_dir():
        dst = dst / src.name
        dst, error = _guard(dst)
        if error:
            return None, error
    if dst.exists():
        return None, _error("exists")
    return dst, None


def copy(src: str, dst: str, overwrite: bool = False) -> dict[str, Any]:
    source, error = _guard(src)
    if error:
        return error
    if not source.exists():
        return _error("not_found")
    tree_error = _validate_tree(source)
    if tree_error:
        return tree_error
    target, error = _destination(source, dst)
    if error and not (overwrite and error["reason"] == "exists"):
        return error
    if overwrite:
        target, error = _guard(dst)
        if error:
            return error
        if target.is_dir():
            target = target / source.name
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        if source.is_dir():
            source_hash = _tree_hash(source)
            shutil.copytree(source, target, dirs_exist_ok=overwrite)
            verified = target.is_dir() and _tree_hash(target) == source_hash
        else:
            shutil.copy2(source, target)
            verified = target.is_file() and _hash(source) == _hash(target)
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"hash": _hash(target) if target.is_file() else None}}
    except OSError as error:
        return _error("copy_failed", message=str(error))


def move(src: str, dst: str, overwrite: bool = False) -> dict[str, Any]:
    source, error = _guard(src)
    if error:
        return error
    if not source.exists():
        return _error("not_found")
    tree_error = _validate_tree(source)
    if tree_error:
        return tree_error
    target, error = _destination(source, dst)
    if error and not (overwrite and error["reason"] == "exists"):
        return error
    if overwrite:
        target, error = _guard(dst)
        if error:
            return error
        if target.is_dir():
            target = target / source.name
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        source_hash = _hash(source) if source.is_file() else _tree_hash(source) if source.is_dir() else None
        if target.exists() and overwrite:
            from send2trash import send2trash
            send2trash(str(target))
        shutil.move(str(source), str(target))
        verified = target.exists() and not source.exists()
        if verified and source_hash and target.is_file():
            verified = _hash(target) == source_hash
        elif verified and source_hash and target.is_dir():
            verified = _tree_hash(target) == source_hash
        return {"ok": verified, "success": verified, "verified": verified, "path": str(target), "details": {"source_absent": not source.exists()}}
    except OSError as error:
        return _error("move_failed", message=str(error))


def delete(path: str) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if not target.exists():
        return _error("not_found")
    tree_error = _validate_tree(target)
    if tree_error:
        return tree_error
    try:
        from send2trash import send2trash
        send2trash(str(target))
        verified = not target.exists()
        return {"ok": verified, "success": verified, "verified": verified, "details": {"path": str(target), "recycled": verified}}
    except (OSError, ImportError) as error:
        logger.warning("Unable to send path to Recycle Bin: %s", target, exc_info=True)
        return _error("recycle_failed", message=str(error))


def open_file(path: str) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if not target.is_file():
        return _error("not_found")
    try:
        if os.name != "nt":
            return _error("windows_only")
        os.startfile(str(target))
        return {"ok": True, "success": True, "verified": True, "path": str(target)}
    except OSError as error:
        return _error("open_failed", message=str(error))


def open_in_explorer(path: str) -> dict[str, Any]:
    target, error = _guard(path)
    if error:
        return error
    if not target.exists():
        return _error("not_found")
    try:
        if os.name != "nt":
            return _error("windows_only")
        if target.is_dir():
            subprocess.Popen(["explorer.exe", str(target)], close_fds=True)
        else:
            subprocess.Popen(["explorer.exe", "/select,", str(target)], close_fds=True)
        return {"ok": True, "success": True, "verified": True, "path": str(target)}
    except OSError as error:
        return _error("open_failed", message=str(error))


def plan_bulk(op: str, root: str, pattern: str, destination: str = "") -> dict[str, Any]:
    if op not in {"copy", "move", "delete", "rename"}:
        return _error("invalid_bulk_operation")
    base, error = _guard(root)
    if error:
        return error
    if not base.is_dir():
        return _error("not_a_directory")
    entries = []
    for path in _walk_files(base):
        if fnmatch.fnmatch(path.name.casefold(), pattern.casefold()):
            entries.append(path)
            if len(entries) > _MAX_BULK_ITEMS:
                return _error("limit_exceeded", limit=_MAX_BULK_ITEMS)
    planned = []
    for src in entries:
        if op == "delete":
            planned.append({"op": "delete", "src": str(src), "dst": ""})
        elif op == "rename":
            if not destination:
                return _error("missing_rename_template")
            try:
                name = destination.format(name=src.name, stem=src.stem, ext=src.suffix)
            except (KeyError, ValueError, IndexError):
                return _error("invalid_rename_template")
            dst = src.parent / name
            resolved_dst, dst_error = _guard(dst)
            if dst_error:
                return dst_error
            if resolved_dst.exists():
                return _error("exists", path=str(resolved_dst))
            planned.append({"op": "rename", "src": str(src), "dst": str(resolved_dst)})
        else:
            dest_base, dest_error = _guard(destination)
            if dest_error:
                return dest_error
            dst = dest_base / src.name
            resolved_dst, dst_error = _guard(dst)
            if dst_error:
                return dst_error
            if resolved_dst.exists():
                return _error("exists", path=str(resolved_dst))
            planned.append({"op": op, "src": str(src), "dst": str(resolved_dst)})
    return {"ok": True, "success": True, "verified": True, "planned": planned, "count": len(planned)}


def execute_bulk(plan: dict[str, Any]) -> dict[str, Any]:
    outcomes = []
    for item in plan.get("planned", []):
        op = item["op"]
        result = delete(item["src"]) if op == "delete" else copy(item["src"], item["dst"]) if op == "copy" else move(item["src"], item["dst"])
        outcomes.append(result)
        if not result.get("verified"):
            return {"ok": False, "success": False, "verified": False, "results": outcomes}
    return {"ok": True, "success": True, "verified": True, "results": outcomes, "count": len(outcomes)}
